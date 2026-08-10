"""
Tests for the greenfield XLSX exporter.

Builds a workbook from the Seatrium reference numbers and asserts:
  - the file loads back via openpyxl
  - all 6 sheets exist
  - formulas are present (not pre-computed literals)
  - toggle cells for shared services are editable inputs

Run directly (no conftest/.env dependency):
    python3 tests/unit/test_exporters_greenfield.py
"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import openpyxl

_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_ROOT))

from data_agents.cost_app.databricks.exporters_greenfield import (  # noqa: E402
    build_greenfield_xlsx,
)


def _sample_workbook() -> BytesIO:
    databricks_lines = [
        {"code": "W01", "name": "Batch ETL (Jobs Photon)", "monthly_usd": 571.54},
        {"code": "W02", "name": "Lakeflow Connect", "monthly_usd": 900.00},
        {"code": "W03", "name": "SQL Warehouse Serverless", "monthly_usd": 1858.56},
        {"code": "W04", "name": "Notebook", "monthly_usd": 316.10},
        {"code": "W05", "name": "Monitoring", "monthly_usd": 372.48},
        {"code": "W06", "name": "Genie (net free allowance)", "monthly_usd": 0.0},
        {"code": "W07", "name": "AI Search", "monthly_usd": 128.48},
        {"code": "W08", "name": "FM Claude Opus tokens", "monthly_usd": 472.50},
    ]
    azure_infra_per_env = [
        {
            "ref": "S01",
            "resource": "NAT Gateway (hours)",
            "qty": 1,
            "hours_or_gb": 730,
            "unit_price": 0.045,
        },
        {
            "ref": "S07",
            "resource": "Log Analytics ingest 200GB",
            "qty": 1,
            "hours_or_gb": 200,
            "unit_price": 2.99,
        },
    ]
    azure_infra_shared = [
        {
            "ref": "S14",
            "resource": "Azure Firewall Standard",
            "enabled": 1,
            "qty": 1,
            "hours_or_gb": 730,
            "unit_price": 1.25,
        },
        {
            "ref": "S16",
            "resource": "Defender for Servers P2",
            "enabled": 1,
            "qty": 16,
            "hours_or_gb": 730,
            "unit_price": 0.02,
        },
    ]
    storage_line = {
        "name": "ADLS Gen2 Hot 50TB",
        "gb": 50000,
        "price_per_gb": 0.0208,
        "monthly_usd": 1040.0,
    }

    return build_greenfield_xlsx(
        client="Seatrium",
        region="southeastasia",
        tier="Premium",
        currency="USD",
        databricks_lines=databricks_lines,
        azure_infra_per_env=azure_infra_per_env,
        azure_infra_shared=azure_infra_shared,
        storage_line=storage_line,
        num_envs=4,
    )


def test_workbook_loads_and_has_all_sheets():
    buf = _sample_workbook()
    wb = openpyxl.load_workbook(buf)
    expected = [
        "1. Cover",
        "2. Databricks",
        "3. Azure Infra",
        "4. Rollup",
        "5. Summary",
        "6. Sources",
    ]
    assert wb.sheetnames == expected


def test_databricks_sheet_has_formula_total():
    wb = openpyxl.load_workbook(_sample_workbook())
    ws = wb["2. Databricks"]
    # Some cell in column C must be a SUM formula (not a literal).
    formulas = [
        ws.cell(row=r, column=3).value
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=3).value, str)
        and str(ws.cell(row=r, column=3).value).startswith("=SUM")
    ]
    assert formulas, "expected at least one =SUM formula in Databricks totals"


def test_azure_infra_toggle_cells_are_editable_inputs():
    wb = openpyxl.load_workbook(_sample_workbook())
    ws = wb["3. Azure Infra"]
    # Find the shared-services 'On?' toggles (value 0 or 1, yellow fill).
    toggles = []
    for row in ws.iter_rows():
        for cell in row:
            if (
                cell.value in (0, 1)
                and cell.fill
                and cell.fill.fgColor.rgb
                and "FFF3B0" in str(cell.fill.fgColor.rgb)
            ):
                toggles.append(cell.coordinate)
    assert toggles, "expected editable ON/OFF toggle cells with yellow fill"


def test_infra_cost_formula_multiplies_toggle():
    """A shared-service cost formula must include the toggle so OFF→0 zeroes it."""
    wb = openpyxl.load_workbook(_sample_workbook())
    ws = wb["3. Azure Infra"]
    cost_formulas = [
        cell.value
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=C") and "*" in cell.value
    ]
    # At least one shared cost formula multiplies 4 factors (toggle × qty × h × price)
    assert any(f.count("*") >= 3 for f in cost_formulas)


def test_rollup_grand_total_is_formula():
    wb = openpyxl.load_workbook(_sample_workbook())
    ws = wb["4. Rollup"]
    sums = [
        cell.value
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=SUM")
    ]
    assert sums, "rollup grand total must be a SUM formula"


def test_rollup_storage_row_is_wired():
    """The storage row in the rollup must reference the Databricks storage cell,
    not be blank (regression for the _cell wiring bug)."""
    wb = openpyxl.load_workbook(_sample_workbook())
    ws = wb["4. Rollup"]
    storage_refs = [
        cell.value
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "'2. Databricks'" in cell.value
    ]
    assert storage_refs, "rollup must reference the storage cell on the Databricks sheet"


if __name__ == "__main__":
    passed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"  ✅ {name}")
            passed += 1
    print(f"\n{passed} greenfield exporter tests passed.")
