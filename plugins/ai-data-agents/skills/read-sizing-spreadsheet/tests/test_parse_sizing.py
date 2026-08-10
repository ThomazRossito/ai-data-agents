"""
Unit tests for the read-sizing-spreadsheet skill.

Builds a synthetic Seatrium-style sizing workbook in-memory (openpyxl) so the
test has no external file dependency, then asserts the parser + interpreter
extract the expected scenario. Also includes a robustness test for the prose
note that previously fooled header detection.

Run: pytest skills/finops/read-sizing-spreadsheet/tests/test_parse_sizing.py -q
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

# Make the skill importable regardless of cwd.
_SKILL_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_SKILL_DIR))

from parse_sizing import (  # noqa: E402
    parse_sizing,
    sizing_to_scenario,
    _extract_storage_gb,
    _extract_region,
)


def _build_seatrium_like(path: Path) -> None:
    """Recreate the structure of the real Seatrium sizing sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sizing Assumptions"

    ws["A1"] = "Seatrium × Databricks — Platform Sizing Assumptions"
    ws["A2"] = (
        "Inputs used to size the Databricks platform estimate. Commercial rates "
        "are intentionally excluded."
    )
    # Prose note containing the word 'assumptions' — must NOT be detected as header.
    ws["A3"] = (
        "Sources: Seatrium-provided requirements and customer inputs · account "
        "assumptions where marked. T-shirt-band, indicative sizing."
    )
    # Real header on row 4.
    headers = ["#", "Assumption / driver", "Value used in estimate", "Basis / source", "Notes for CI&T"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)

    rows = [
        ("A3", "Environments", "4 total (1 PROD + 2 non-prod + 1 Sandbox)", "Account assumption", ""),
        ("A4", "Region / residency", "Single-region Singapore", "Customer input", ""),
        ("B1", "Total data at rest", "Band L (10–100 TB) → 50,000 GB modelled", "Customer", "Raw + curated"),
        ("B2", "Annual data growth", "20% per year", "Customer requirement", ""),
        ("B3", "Batch ETL volume", "Band S (<50 GB/day) → 25 GB/day modelled", "Customer", ""),
        ("B4", "Streaming / real-time", "None — batch only", "Customer input", ""),
        ("C1", "Source systems", "10 systems (incl. Ivalua)", "Customer requirement", ""),
        ("D1", "SQL / BI workload", "70 Daily interactive users, assume all is using it.", "Customer input", ""),
        ("D2", "Dashboards", "estimated around 10 dashboard", "Account assumption", ""),
        ("E1", "Notebook users (DS/DE)", "10 users modelled", "Account assumption", ""),
        ("F1", "Genie Q&A volume", "5,000 questions/mo", "Account assumption", ""),
        ("F2", "Genie user base", "70 users; 5% power users at ~2× volume", "Customer input", ""),
        ("G1", "Monitoring coverage", "20 core tables profiled weekly (×4 runs/mo)", "Account", ""),
        ("H1", "ML training", "None modelled at this stage", "Account decision", ""),
        ("I2", "LLM", "Proprietary frontier model (Claude Opus class) via Foundation Model API", "Account", ""),
        ("I3", "GenAI question volume", "5,000 questions/mo", "Account assumption", ""),
        ("I5", "RAG / retrieval", "Yes — Databricks AI Search vector retrieval endpoint sized", "Use-case", ""),
        ("K2", "Lakebase / OLTP serving", "Not sized", "Use case to confirm", ""),
        ("K3", "DR / HA targets", "None stated", "Customer input", ""),
        ("K4", "Historical data migration", "None in scope", "Customer input", ""),
    ]
    r = 5
    for row in rows:
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)
        r += 1
    wb.save(path)


@pytest.fixture()
def seatrium_xlsx(tmp_path: Path) -> Path:
    p = tmp_path / "seatrium_sizing.xlsx"
    _build_seatrium_like(p)
    return p


def test_parse_finds_correct_header_row(seatrium_xlsx: Path):
    parsed = parse_sizing(seatrium_xlsx)
    # Header must be row 4, not row 3 (the prose note with 'assumptions').
    assert parsed["_sheet"] == "Sizing Assumptions"
    # Must have captured data rows, not the prose note as a field.
    assert len(parsed["raw_rows"]) >= 18


def test_scenario_core_fields(seatrium_xlsx: Path):
    scenario = sizing_to_scenario(parse_sizing(seatrium_xlsx))
    assert scenario["region"] == "southeastasia"
    assert scenario["num_envs"] == 4
    assert scenario["storage_gb"] == 50000  # explicit GB, not the 10-TB band
    assert scenario["num_sql_users"] == 70  # matched via "SQL / BI workload"
    assert scenario["num_notebook_users"] == 10
    assert scenario["genie_questions_per_month"] == 5000
    assert scenario["annual_growth_pct"] == 20.0
    assert "Claude Opus" in scenario["genai_llm"]
    assert scenario["needs_confirmation"] == []


def test_scenario_workloads_and_exclusions(seatrium_xlsx: Path):
    scenario = sizing_to_scenario(parse_sizing(seatrium_xlsx))
    codes = {w["code"] for w in scenario["workloads"]}
    assert {"ETL_BATCH", "SQL_WAREHOUSE", "NOTEBOOK", "GENIE", "FM_LLM"} <= codes
    assert scenario["streaming_required"] is False
    assert "streaming" in scenario["excluded"]
    assert "lakebase" in scenario["excluded"]


def test_storage_prefers_explicit_gb():
    assert _extract_storage_gb("Band L (10–100 TB) → 50,000 GB modelled") == 50000
    assert _extract_storage_gb("~10 TB estate") == 10000
    assert _extract_storage_gb("nonsense") is None


def test_region_mapping():
    assert _extract_region("Single-region Singapore") == "southeastasia"
    assert _extract_region("Brazil South") == "brazilsouth"
    assert _extract_region("Mars colony") is None


def test_missing_required_field_flags_confirmation(tmp_path: Path):
    """A sheet without a region should surface it in needs_confirmation, not guess."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sizing"
    for i, h in enumerate(["#", "Assumption", "Value used", "Basis"], 1):
        ws.cell(row=1, column=i, value=h)
    ws.append(["A1", "Batch ETL volume", "25 GB/day", "Customer"])
    p = tmp_path / "partial.xlsx"
    wb.save(p)

    scenario = sizing_to_scenario(parse_sizing(p))
    assert scenario["region"] is None
    assert any("cloud_region" in c for c in scenario["needs_confirmation"])
